import ipaddress
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font

# ============ 配置区 ============
BASE_NETWORK = "37.55.0.0/16"   # 地址池
VLAN_START = 100                # VLAN 起始编号

# 前缀 → 可用 IP 数（总地址 - 网络 - 广播 - 网关）
PREFIX_CAPACITY = {
    30: 1,
    29: 5,
    28: 13,
    27: 29,
    26: 61,
    25: 125,
    24: 253,
    23: 509,
    22: 1021,
    21: 2045,
    20: 4093,
}

def pick_prefix(need: int) -> int:
    """选择能容纳 need 个可用 IP 的最小前缀"""
    for prefix in sorted(PREFIX_CAPACITY.keys(), reverse=True):
        if PREFIX_CAPACITY[prefix] >= need:
            return prefix
    raise ValueError(f"需求 {need} 超过单个 /20 子网容量")

def calculate_ip_details(ip_demand, current_int, base_net):
    """根据需求与当前地址指针，返回一个子网的各项信息，及新的指针位置"""
    prefix = pick_prefix(ip_demand)
    size = 2 ** (32 - prefix)

    # 对齐到子网边界
    if current_int % size != 0:
        current_int = ((current_int // size) + 1) * size

    network = ipaddress.ip_network((ipaddress.ip_address(current_int), prefix))

    if int(network.broadcast_address) > int(base_net.broadcast_address):
        raise ValueError("地址池已耗尽")

    broadcast = network.broadcast_address
    gateway = ipaddress.ip_address(int(broadcast) - 1)

    available = [ip for ip in network.hosts() if ip != gateway]
    if len(available) == 1:
        avail_str = str(available[0])
    else:
        avail_str = f"{available[0]} ~ {available[-1]}（共 {len(available)} 个）"

    return {
        "subnet": str(network),
        "netmask": str(network.netmask),
        "gateway": str(gateway),
        "available": avail_str,
        "next_int": int(broadcast) + 1,
    }

def main():
    # 创建 Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "IP Address Planning"

    headers = ["IP需求", "建议网段", "子网掩码", "网关", "可用IP地址", "VLAN"]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = Font(bold=True)

    # 输入需求列表（可改成从别处读取）
    ip_demands = [1, 5, 10, 20, 50, 100]

    base_net = ipaddress.ip_network(BASE_NETWORK)
    current_int = int(base_net.network_address)

    for idx, demand in enumerate(ip_demands):
        row = idx + 2
        ws.cell(row=row, column=1, value=demand)

        try:
            result = calculate_ip_details(demand, current_int, base_net)
        except ValueError as e:
            ws.cell(row=row, column=2, value=f"错误：{e}")
            continue

        ws.cell(row=row, column=2, value=result["subnet"])
        ws.cell(row=row, column=3, value=result["netmask"])
        ws.cell(row=row, column=4, value=result["gateway"])
        ws.cell(row=row, column=5, value=result["available"])
        ws.cell(row=row, column=6, value=VLAN_START + idx)

        current_int = result["next_int"]

    # 调整列宽，便于查看
    for col, width in zip("ABCDEF", [10, 18, 18, 18, 40, 8]):
        ws.column_dimensions[col].width = width

    wb.save("ip_address_planning.xlsx")
    print("已生成 ip_address_planning.xlsx")

if __name__ == "__main__":
    main()